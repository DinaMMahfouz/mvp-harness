"""Build the 4-sheet harness_issue_register.xlsx from live DB data.

Sheets: Findings, Run Summary, Remediation Tracker, Before/After.
"""
from __future__ import annotations

import io
import uuid

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlmodel import Session

from app.models import Application
from app.reports import queries

HEADER_FONT = Font(bold=True)


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT


def _autosize(ws: Worksheet) -> None:
    for column_cells in ws.columns:
        length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 60)


def build_findings_sheet(wb: Workbook, session: Session, application: Application) -> dict:
    ws = wb.active
    ws.title = "Findings"
    headers = [
        "Finding ID",
        "Category",
        "Severity",
        "Title",
        "Status",
        "Root Cause",
        "Recommendation",
        "First Seen Run",
        "Last Seen Run",
        "Created At",
        "Resolved At",
    ]
    _write_header(ws, headers)
    findings = queries.findings_for_application(session, application.id)
    for f in findings:
        ws.append(
            [
                str(f.id),
                f.category.value,
                f.severity.value,
                f.title,
                f.status.value,
                f.root_cause,
                f.recommendation,
                str(f.first_seen_run_id),
                str(f.last_seen_run_id),
                f.created_at.isoformat(),
                f.resolved_at.isoformat() if f.resolved_at else "",
            ]
        )
    _autosize(ws)
    return {"row_count": len(findings)}


def build_run_summary_sheet(wb: Workbook, session: Session, application: Application) -> dict:
    ws = wb.create_sheet("Run Summary")
    headers = [
        "Run ID",
        "Run Type",
        "Status",
        "Triggered By",
        "Started At",
        "Completed At",
        "Assurance Score",
        "Result Count",
    ]
    _write_header(ws, headers)
    runs = queries.runs_for_application(session, application.id)
    for r in runs:
        results = queries.results_for_run(session, r.id)
        ws.append(
            [
                str(r.id),
                r.run_type.value,
                r.status.value,
                r.triggered_by,
                r.started_at.isoformat() if r.started_at else "",
                r.completed_at.isoformat() if r.completed_at else "",
                r.assurance_score if r.assurance_score is not None else "",
                len(results),
            ]
        )
    _autosize(ws)
    return {"row_count": len(runs)}


def build_remediation_tracker_sheet(wb: Workbook, session: Session, application: Application) -> dict:
    ws = wb.create_sheet("Remediation Tracker")
    headers = ["Finding ID", "Status Change", "Note", "Changed By", "Changed At"]
    _write_header(ws, headers)
    findings = queries.findings_for_application(session, application.id)
    row_count = 0
    for f in findings:
        for rem in queries.remediations_for_finding(session, f.id):
            ws.append(
                [str(f.id), rem.status.value, rem.note or "", rem.changed_by, rem.changed_at.isoformat()]
            )
            row_count += 1
    _autosize(ws)
    return {"row_count": row_count}


def build_before_after_sheet(wb: Workbook, session: Session, application: Application) -> dict:
    ws = wb.create_sheet("Before-After")
    headers = [
        "Comparison ID",
        "Baseline Run",
        "Retest Run",
        "Test Case ID",
        "Classification",
        "Baseline Result",
        "Retest Result",
        "Created At",
    ]
    _write_header(ws, headers)
    comparisons = queries.comparisons_for_application(session, application.id)
    row_count = 0
    for c in comparisons:
        for detail in c.details:
            ws.append(
                [
                    str(c.id),
                    str(c.baseline_run_id),
                    str(c.retest_run_id),
                    detail.get("test_case_id", ""),
                    detail.get("classification", ""),
                    detail.get("baseline_result", ""),
                    detail.get("retest_result", ""),
                    c.created_at.isoformat(),
                ]
            )
            row_count += 1
    _autosize(ws)
    return {"row_count": row_count}


def build_workbook(session: Session, application: Application) -> tuple[io.BytesIO, dict]:
    wb = Workbook()
    summary = {
        "application_id": str(application.id),
        "application_name": application.name,
        "findings": build_findings_sheet(wb, session, application),
        "run_summary": build_run_summary_sheet(wb, session, application),
        "remediation_tracker": build_remediation_tracker_sheet(wb, session, application),
        "before_after": build_before_after_sheet(wb, session, application),
    }
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, summary
